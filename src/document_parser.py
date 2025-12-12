"""Document parser module for handling multiple file formats."""

import os
from pathlib import Path
from typing import Dict, Optional
import PyPDF2
import pdfplumber
from docx import Document
import logging

logger = logging.getLogger(__name__)


class DocumentParser:
    """Parse various document formats and extract text content."""
    
    SUPPORTED_FORMATS = ['.pdf', '.docx', '.txt', '.md']
    
    def __init__(self):
        """Initialize the document parser."""
        self.parsers = {
            '.pdf': self._parse_pdf,
            '.docx': self._parse_docx,
            '.txt': self._parse_txt,
            '.md': self._parse_txt,
        }
    
    def parse(self, file_path: str) -> Dict[str, any]:
        """
        Parse a document and extract its content.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            Dictionary containing:
                - content: Extracted text content
                - file_name: Name of the file
                - file_type: Type/extension of the file
                - metadata: Additional metadata if available
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = path.suffix.lower()
        
        if file_ext not in self.SUPPORTED_FORMATS:
            raise ValueError(
                f"Unsupported file format: {file_ext}. "
                f"Supported formats: {', '.join(self.SUPPORTED_FORMATS)}"
            )
        
        parser_func = self.parsers[file_ext]
        content = parser_func(file_path)
        
        return {
            'content': content,
            'file_name': path.name,
            'file_type': file_ext,
            'metadata': self._extract_metadata(file_path, file_ext)
        }
    
    def _parse_pdf(self, file_path: str) -> str:
        """Parse PDF file and extract text."""
        text_content = []
        
        try:
            # Try pdfplumber first (better text extraction)
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_content.append(text)
        except Exception as e:
            logger.warning(f"pdfplumber failed, trying PyPDF2: {e}")
            # Fallback to PyPDF2
            try:
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page in pdf_reader.pages:
                        text = page.extract_text()
                        if text:
                            text_content.append(text)
            except Exception as e2:
                logger.error(f"Both PDF parsers failed: {e2}")
                raise ValueError(f"Could not parse PDF file: {file_path}")
        
        return '\n\n'.join(text_content)
    
    def _parse_docx(self, file_path: str) -> str:
        """Parse DOCX file and extract text."""
        try:
            doc = Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            return '\n\n'.join(paragraphs)
        except Exception as e:
            logger.error(f"Failed to parse DOCX: {e}")
            raise ValueError(f"Could not parse DOCX file: {file_path}")
    
    def _parse_txt(self, file_path: str) -> str:
        """Parse text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                logger.error(f"Failed to parse text file: {e}")
                raise ValueError(f"Could not parse text file: {file_path}")
    
    def _extract_metadata(self, file_path: str, file_type: str) -> Dict:
        """Extract metadata from document."""
        metadata = {
            'file_size': os.path.getsize(file_path),
        }
        
        if file_type == '.pdf':
            try:
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    metadata['page_count'] = len(pdf_reader.pages)
                    if pdf_reader.metadata:
                        metadata['title'] = pdf_reader.metadata.get('/Title', '')
                        metadata['author'] = pdf_reader.metadata.get('/Author', '')
            except Exception as e:
                logger.warning(f"Could not extract PDF metadata: {e}")
        
        elif file_type == '.docx':
            try:
                doc = Document(file_path)
                metadata['paragraph_count'] = len(doc.paragraphs)
                core_props = doc.core_properties
                metadata['title'] = core_props.title or ''
                metadata['author'] = core_props.author or ''
            except Exception as e:
                logger.warning(f"Could not extract DOCX metadata: {e}")
        
        return metadata
