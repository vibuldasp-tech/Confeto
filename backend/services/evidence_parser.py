"""
Evidence Parser - Parse supporting evidence files
Extract text from PDFs, DOCX, XLSX, PPTX, images
Build searchable index of evidence content
"""

from pathlib import Path
from typing import Dict, Any, List
import re


class EvidenceParser:
    def __init__(self):
        self.parsers = {
            '.pdf': self._parse_pdf,
            '.docx': self._parse_docx,
            '.xlsx': self._parse_xlsx,
            '.csv': self._parse_csv,
            '.pptx': self._parse_pptx,
            '.png': self._parse_image,
            '.jpg': self._parse_image,
            '.jpeg': self._parse_image
        }
    
    def parse_file(self, file_path: str) -> Dict[str, Any]:
        """Parse evidence file and extract content"""
        path = Path(file_path)
        ext = path.suffix.lower()
        
        parser = self.parsers.get(ext)
        if not parser:
            return {
                'filename': path.name,
                'type': ext,
                'error': f'Unsupported file type: {ext}',
                'text': '',
                'pages': []
            }
        
        try:
            return parser(file_path)
        except Exception as e:
            return {
                'filename': path.name,
                'type': ext,
                'error': str(e),
                'text': '',
                'pages': []
            }
    
    def _parse_pdf(self, file_path: str) -> Dict[str, Any]:
        """Parse PDF file"""
        try:
            import fitz  # PyMuPDF
            
            doc = fitz.open(file_path)
            pages = []
            full_text = []
            
            for page_num, page in enumerate(doc):
                text = page.get_text()
                pages.append({
                    'page': page_num + 1,
                    'text': text,
                    'word_count': len(text.split())
                })
                full_text.append(text)
            
            doc.close()
            
            return {
                'filename': Path(file_path).name,
                'type': 'pdf',
                'text': '\n\n'.join(full_text),
                'pages': pages,
                'page_count': len(pages)
            }
        except Exception as e:
            raise Exception(f"Error parsing PDF: {str(e)}")
    
    def _parse_docx(self, file_path: str) -> Dict[str, Any]:
        """Parse DOCX file"""
        try:
            from docx import Document
            
            doc = Document(file_path)
            paragraphs = []
            
            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if text:
                    paragraphs.append({
                        'paragraph': i + 1,
                        'text': text
                    })
            
            # Extract tables
            tables = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)
            
            full_text = '\n\n'.join([p['text'] for p in paragraphs])
            
            return {
                'filename': Path(file_path).name,
                'type': 'docx',
                'text': full_text,
                'paragraphs': paragraphs,
                'tables': tables,
                'paragraph_count': len(paragraphs)
            }
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
    
    def _parse_xlsx(self, file_path: str) -> Dict[str, Any]:
        """Parse Excel file"""
        try:
            import pandas as pd
            
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            sheets = []
            full_text = []
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                
                # Convert to text
                sheet_text = df.to_string()
                sheets.append({
                    'sheet_name': sheet_name,
                    'rows': len(df),
                    'columns': len(df.columns),
                    'text': sheet_text
                })
                full_text.append(f"Sheet: {sheet_name}\n{sheet_text}")
            
            return {
                'filename': Path(file_path).name,
                'type': 'xlsx',
                'text': '\n\n'.join(full_text),
                'sheets': sheets,
                'sheet_count': len(sheets)
            }
        except Exception as e:
            raise Exception(f"Error parsing XLSX: {str(e)}")
    
    def _parse_csv(self, file_path: str) -> Dict[str, Any]:
        """Parse CSV file"""
        try:
            import pandas as pd
            
            df = pd.read_csv(file_path)
            text = df.to_string()
            
            return {
                'filename': Path(file_path).name,
                'type': 'csv',
                'text': text,
                'rows': len(df),
                'columns': len(df.columns)
            }
        except Exception as e:
            raise Exception(f"Error parsing CSV: {str(e)}")
    
    def _parse_pptx(self, file_path: str) -> Dict[str, Any]:
        """Parse PowerPoint file"""
        try:
            from pptx import Presentation
            
            prs = Presentation(file_path)
            slides = []
            full_text = []
            
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)
                
                slide_content = '\n'.join(slide_text)
                slides.append({
                    'slide': i + 1,
                    'text': slide_content
                })
                full_text.append(slide_content)
            
            return {
                'filename': Path(file_path).name,
                'type': 'pptx',
                'text': '\n\n'.join(full_text),
                'slides': slides,
                'slide_count': len(slides)
            }
        except Exception as e:
            raise Exception(f"Error parsing PPTX: {str(e)}")
    
    def _parse_image(self, file_path: str) -> Dict[str, Any]:
        """Parse image file (OCR)"""
        try:
            import pytesseract
            from PIL import Image
            
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            
            return {
                'filename': Path(file_path).name,
                'type': 'image',
                'text': text,
                'ocr': True
            }
        except Exception as e:
            # Return empty if OCR fails
            return {
                'filename': Path(file_path).name,
                'type': 'image',
                'text': '',
                'error': f'OCR failed: {str(e)}'
            }
    
    def search_text(self, parsed_files: List[Dict[str, Any]], query: str, threshold: float = 0.3) -> List[Dict[str, Any]]:
        """Search for query text across all parsed files"""
        results = []
        query_lower = query.lower()
        
        for file_data in parsed_files:
            text = file_data.get('text', '').lower()
            
            if query_lower in text:
                # Find context around match
                match_start = text.find(query_lower)
                context_start = max(0, match_start - 100)
                context_end = min(len(text), match_start + len(query) + 100)
                context = text[context_start:context_end]
                
                results.append({
                    'file_id': file_data.get('file_id'),
                    'filename': file_data.get('filename'),
                    'match': query,
                    'context': context,
                    'confidence': 0.8
                })
        
        return results
